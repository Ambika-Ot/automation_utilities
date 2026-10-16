import os
import re
import shutil
import sys
import time
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from xml.dom import minidom

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def clean_tag(name, position):
    # xml tags cant have spaces or symbols and cant start with a number
    name = str(name).strip() if name is not None else ""
    name = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if name == "":
        name = "column" + str(position)
    if name[0].isdigit():
        name = "_" + name
    return name


def convert(excel_file, xml_file, sheet_name=None):
    wb = load_workbook(
        excel_file, data_only=True
    )  # data_only gives values instead of formulas

    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) == 0:
        print("sheet is empty")
        return False

    # first row = headers
    headers = []
    for i, h in enumerate(rows[0]):
        headers.append(clean_tag(h, i + 1))

    root = ET.Element(clean_tag(ws.title, 0))

    count = 0
    for row in rows[1:]:
        # skip completely empty rows
        if all(cell is None for cell in row):
            continue

        record = ET.SubElement(root, "row")
        for header, value in zip(headers, row):
            field = ET.SubElement(record, header)
            if value is not None:
                field.text = str(value)
        count += 1

    # make it look nice with indentation
    text = ET.tostring(root, encoding="utf-8")
    pretty = minidom.parseString(text).toprettyxml(indent="  ", encoding="utf-8")

    with open(xml_file, "wb") as f:
        f.write(pretty)

    print("converted", count, "rows ->", xml_file)
    return True


CHECK_EVERY = 3  # seconds between checking the folder


def file_is_ready(path):
    # a file that is still being copied keeps growing, so wait and compare the size
    size1 = os.path.getsize(path)
    time.sleep(1)
    size2 = os.path.getsize(path)
    return size1 == size2


def move_to(path, folder):
    dest = os.path.join(folder, os.path.basename(path))
    if os.path.exists(dest):
        # same name already there, add the time so nothing gets overwritten
        base, ext = os.path.splitext(dest)
        dest = base + "_" + datetime.now().strftime("%H%M%S") + ext
    shutil.move(path, dest)


def watch(folder, out_folder=None):
    if out_folder is None:
        out_folder = os.path.join(folder, "xml_output")
    done_folder = os.path.join(folder, "processed")
    failed_folder = os.path.join(folder, "failed")

    for f in (folder, out_folder, done_folder, failed_folder):
        os.makedirs(f, exist_ok=True)

    print("watching", folder, "- drop excel files in it (Ctrl+C to stop)")

    try:
        while True:
            for name in os.listdir(folder):
                path = os.path.join(folder, name)

                if not os.path.isfile(path):
                    continue
                if name.startswith("~$"):
                    continue  # temp file excel makes while a file is open
                if not name.lower().endswith((".xlsx", ".xlsm")):
                    continue
                if not file_is_ready(path):
                    continue

                xml_path = os.path.join(out_folder, os.path.splitext(name)[0] + ".xml")
                print("found", name)

                try:
                    ok = convert(path, xml_path)
                    move_to(path, done_folder if ok else failed_folder)
                except PermissionError:
                    print("file is open somewhere, will try again:", name)
                except (
                    OSError,
                    ValueError,
                    KeyError,
                    ET.ParseError,
                    zipfile.BadZipFile,
                    InvalidFileException,
                ) as e:
                    print("failed:", name, "-", e)
                    move_to(path, failed_folder)

            time.sleep(CHECK_EVERY)
    except KeyboardInterrupt:
        print("\nstopped")


def main():
    if len(sys.argv) < 2:
        print("usage: python excel_to_xml.py data.xlsx [output.xml] [sheet name]")
        print("       python excel_to_xml.py --watch folder [output folder]")
        return

    if sys.argv[1] == "--watch":
        if len(sys.argv) < 3:
            print("give the folder to watch")
            return
        out_folder = sys.argv[3] if len(sys.argv) > 3 else None
        watch(sys.argv[2], out_folder)
        return

    excel_file = sys.argv[1]
    xml_file = (
        sys.argv[2] if len(sys.argv) > 2 else excel_file.rsplit(".", 1)[0] + ".xml"
    )
    sheet_name = sys.argv[3] if len(sys.argv) > 3 else None

    convert(excel_file, xml_file, sheet_name)


if __name__ == "__main__":
    main()
