import copy
import json
import shutil
import sys
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt

THEME_COLOR_MAP = {
    "dk1": MSO_THEME_COLOR.DARK_1,
    "lt1": MSO_THEME_COLOR.LIGHT_1,
    "dk2": MSO_THEME_COLOR.DARK_2,
    "lt2": MSO_THEME_COLOR.LIGHT_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1,
    "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3,
    "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5,
    "accent6": MSO_THEME_COLOR.ACCENT_6,
}
ACCENT_NAMES = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

THEMES = {
    "corporate_blue": {
        "theme": {
            "heading_font": "Calibri",
            "body_font": "Calibri",
            "colors": {
                "dk2": "#1F4E79",
                "accent1": "#1F4E79",
                "accent2": "#2E75B6",
                "accent3": "#9DC3E6",
                "accent4": "#F28C28",
                "accent5": "#7F7F7F",
                "accent6": "#BDD7EE",
            },
        },
        "text_style": {"title_color": "#1F4E79", "body_color": "#333333"},
        "background_color": "#FFFFFF",
    },
    "dark_mode": {
        "theme": {
            "heading_font": "Segoe UI",
            "body_font": "Segoe UI",
            "colors": {
                "dk2": "#1E1E2E",
                "accent1": "#7AA2F7",
                "accent2": "#BB9AF7",
                "accent3": "#9ECE6A",
                "accent4": "#E0AF68",
                "accent5": "#F7768E",
                "accent6": "#7DCFFF",
            },
        },
        "text_style": {"title_color": "#FFFFFF", "body_color": "#D0D0D8"},
        "background_color": "#1E1E2E",
    },
    "forest_green": {
        "theme": {
            "heading_font": "Georgia",
            "body_font": "Calibri",
            "colors": {
                "dk2": "#1B4332",
                "accent1": "#2D6A4F",
                "accent2": "#40916C",
                "accent3": "#74C69D",
                "accent4": "#B7E4C7",
                "accent5": "#D4A373",
                "accent6": "#8D6E63",
            },
        },
        "text_style": {"title_color": "#1B4332", "body_color": "#2D3A30"},
        "background_color": "#A3D08C",
    },
}


def list_themes():
    print("ready made themes:")
    for name, t in THEMES.items():
        print(
            " ",
            name.ljust(16),
            t["theme"]["heading_font"],
            "/",
            t["theme"]["body_font"],
            "| background",
            t["background_color"],
            "| title",
            t["text_style"]["title_color"],
        )
        colors = t["theme"].get("colors", {})
        accents = ", ".join(
            f"{k}={v}" for k, v in colors.items() if k.startswith("accent")
        )
        if accents:
            print(" " * 4, "accents:", accents)


def build_settings(config):
    theme = {}
    style = {}
    background = None

    name = config.get("preset")
    if name:
        if name not in THEMES:
            print("unknown preset:", name)
            list_themes()
            sys.exit(1)
        preset = copy.deepcopy(THEMES[name])
        theme = preset["theme"]
        style = preset["text_style"]
        background = preset["background_color"]

    for key, value in config.get("theme", {}).items():
        if key == "colors":
            theme.setdefault("colors", {}).update(value)
        else:
            theme[key] = value

    style.update(config.get("text_style", {}))

    if "background_color" in config:
        background = config["background_color"]

    if "heading_font" in theme:
        style.setdefault("title_font", theme["heading_font"])
    if "body_font" in theme:
        style.setdefault("body_font", theme["body_font"])

    return theme, style, background


def hex_to_rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def apply_fill(shape, fill_value):
    """fill_value can be a theme color name ("accent1", "dk2", ...)
    or a literal hex color ("#2D6A4F")."""
    shape.fill.solid()
    if fill_value in THEME_COLOR_MAP:
        shape.fill.fore_color.theme_color = THEME_COLOR_MAP[fill_value]
    else:
        shape.fill.fore_color.rgb = hex_to_rgb(fill_value)


def add_accent_bar(slide, prs):
    """draws one rectangle per accent color (accent1-accent6) across the
    bottom of the slide so every theme accent is actually visible,
    bound symbolically so it stays correct even if the theme is edited
    afterwards by change_theme()."""
    margin = Inches(0.3)
    bar_height = Inches(0.4)
    count = len(ACCENT_NAMES)
    usable_width = prs.slide_width - 2 * margin
    swatch_width = usable_width // count
    top = prs.slide_height - bar_height - margin

    for i, name in enumerate(ACCENT_NAMES):
        left = margin + i * swatch_width
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, swatch_width, bar_height
        )
        apply_fill(box, name)
        box.line.fill.background()  # no outline
        box.shadow.inherit = False
        tf = box.text_frame
        tf.text = name
        tf.paragraphs[0].font.size = Pt(9)


def add_custom_shapes(slide, shapes):
    """lets a slide config list arbitrary shapes, e.g.:
    "shapes": [
      {"shape": "OVAL", "left": 1, "top": 1, "width": 2, "height": 1,
       "fill": "accent2", "text": "Hello"}
    ]
    "fill" accepts either a theme color name (accent1-6, dk1, dk2, lt1, lt2)
    or a literal hex string. left/top/width/height are in inches."""
    for s in shapes:
        shape_name = s.get("shape", "RECTANGLE").upper()
        shape_type = getattr(MSO_SHAPE, shape_name, MSO_SHAPE.RECTANGLE)
        box = slide.shapes.add_shape(
            shape_type,
            Inches(s.get("left", 0.5)),
            Inches(s.get("top", 0.5)),
            Inches(s.get("width", 1.5)),
            Inches(s.get("height", 1.0)),
        )
        if "fill" in s:
            apply_fill(box, s["fill"])
        if "text" in s:
            box.text_frame.text = s["text"]


def all_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)
        else:
            yield sh


def all_paragraphs(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            yield p
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    yield p


def replace_text(prs, replacements):
    for slide in prs.slides:
        for sh in all_shapes(slide.shapes):
            for p in all_paragraphs(sh):
                full = "".join(r.text for r in p.runs)
                changed = full
                for old, new in replacements.items():
                    changed = changed.replace(old, new)
                if changed != full and p.runs:
                    p.runs[0].text = changed
                    for r in p.runs[1:]:
                        r.text = ""


def add_slides(prs, slides, show_accents_default=False):
    for s in slides:
        layout = prs.slide_layouts[s.get("layout", 1)]
        slide = prs.slides.add_slide(layout)

        if "title" in s:
            title = slide.shapes.title
            if title is None:
                title = slide.shapes.add_textbox(
                    Inches(0.75),
                    Inches(0.4),
                    prs.slide_width - Inches(1.5),
                    Inches(0.8),
                )
                title.name = "Generated Title"
            title.text = s["title"]

        if "bullets" in s:
            body = next(
                (
                    placeholder
                    for placeholder in slide.placeholders
                    if placeholder.placeholder_format.type
                    in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
                ),
                None,
            )
            if body is None:
                body = slide.shapes.add_textbox(
                    Inches(0.75),
                    Inches(1.5),
                    prs.slide_width - Inches(1.5),
                    prs.slide_height - Inches(2.0),
                )
            tf = body.text_frame
            tf.text = s["bullets"][0]
            for line in s["bullets"][1:]:
                p = tf.add_paragraph()
                p.text = line

        if "notes" in s:
            slide.notes_slide.notes_text_frame.text = s["notes"]

        if "shapes" in s:
            add_custom_shapes(slide, s["shapes"])

        if s.get("show_accents", show_accents_default):
            add_accent_bar(slide, prs)


def set_background(prs, color):
    for slide in prs.slides:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(color)


def style_text(prs, style):
    for slide in prs.slides:
        for sh in all_shapes(slide.shapes):
            is_title = sh == slide.shapes.title or sh.name == "Generated Title"
            for p in all_paragraphs(sh):
                for run in p.runs:
                    if is_title:
                        font, color, size = (
                            style.get("title_font"),
                            style.get("title_color"),
                            style.get("title_size"),
                        )
                    else:
                        font, color, size = (
                            style.get("body_font"),
                            style.get("body_color"),
                            style.get("body_size"),
                        )
                    if font:
                        run.font.name = font
                    if color:
                        run.font.color.rgb = hex_to_rgb(color)
                    if size:
                        run.font.size = Pt(size)


def change_theme(path, theme):
    tmp = path + ".tmp"
    zin = zipfile.ZipFile(path)
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)

    for item in zin.infolist():
        data = zin.read(item.filename)

        if item.filename == "ppt/theme/theme1.xml":
            root = etree.fromstring(data)

            if "heading_font" in theme:
                tag = root.find(f".//{{{A}}}majorFont/{{{A}}}latin")
                tag.set("typeface", theme["heading_font"])
            if "body_font" in theme:
                tag = root.find(f".//{{{A}}}minorFont/{{{A}}}latin")
                tag.set("typeface", theme["body_font"])

            for name, value in theme.get("colors", {}).items():
                node = root.find(f".//{{{A}}}clrScheme/{{{A}}}{name}")
                if node is None:
                    print("unknown theme color:", name)
                    continue
                for child in list(node):
                    node.remove(child)
                new = etree.SubElement(node, f"{{{A}}}srgbClr")
                new.set("val", value.lstrip("#"))

            data = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        zout.writestr(item, data)

    zin.close()
    zout.close()
    shutil.move(tmp, path)


def run(config):
    prs = Presentation(config["template"])
    out = config.get("output", "output.pptx")
    theme, style, background = build_settings(config)

    if "slides" in config:
        add_slides(
            prs,
            config["slides"],
            show_accents_default=config.get("show_accents", False),
        )
    if "replacements" in config:
        replace_text(prs, config["replacements"])
    if background:
        set_background(prs, background)
    if style:
        style_text(prs, style)

    prs.save(out)

    if theme:
        change_theme(out, theme)

    return out


def main():
    if len(sys.argv) < 2 or "--themes" in sys.argv:
        if "--themes" in sys.argv:
            list_themes()
        else:
            print("usage: python ppt_tool.py config.json")
            print("       python ppt_tool.py --themes")
        return

    config_path = sys.argv[1]
    with open(config_path, encoding="utf-8") as f:
        config = json.load(f)

    out = run(config)
    print("saved", out)


if __name__ == "__main__":
    main()
