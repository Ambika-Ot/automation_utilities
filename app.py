import csv
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import tkinter as tk
import zipfile
from datetime import datetime
from tkinter import colorchooser, filedialog, messagebox, ttk


def try_import(name):
    try:
        return __import__(name), None
    except ImportError as e:
        return None, str(e)


ppt_tool, ppt_err = try_import("ppt_tool")
excel_to_xml, excel_err = try_import("excel_to_xml")
metric_data, metrics_err = try_import("metric_data")

HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")
FONTS = [
    "Calibri",
    "Arial",
    "Segoe UI",
    "Georgia",
    "Cambria",
    "Verdana",
    "Trebuchet MS",
    "Times New Roman",
    "Tahoma",
    "Consolas",
]


def norm_hex(text):
    # "1f4e79" or "#1F4E79" -> "#1F4E79", returns None if it is not a color
    text = text.strip()
    if not text.startswith("#"):
        text = "#" + text
    return text.upper() if HEX_RE.match(text) else None


def open_file(path):
    if hasattr(os, "startfile"):
        os.startfile(path)


class LogBox(ttk.Frame):
    def __init__(self, parent, height=8):
        super().__init__(parent)
        self.text = tk.Text(
            self,
            height=height,
            state="disabled",
            wrap="word",
            bg="#fafafa",
            relief="solid",
            bd=1,
            font=("Consolas", 9),
        )
        bar = ttk.Scrollbar(self, command=self.text.yview)
        self.text.config(yscrollcommand=bar.set)
        self.text.pack(side="left", fill="both", expand=True)
        bar.pack(side="right", fill="y")

    def add(self, line):
        stamp = datetime.now().strftime("%H:%M:%S")
        self.text.config(state="normal")
        self.text.insert("end", stamp + "  " + str(line) + "\n")
        self.text.see("end")
        self.text.config(state="disabled")


class Meter(ttk.Frame):
    def __init__(self, parent, name):
        super().__init__(parent)
        ttk.Label(self, text=name, width=8).pack(side="left")
        self.bar = ttk.Progressbar(self, maximum=100, length=320)
        self.bar.pack(side="left", padx=6)
        self.value = ttk.Label(self, text="-", width=22)
        self.value.pack(side="left")

    def set(self, percent, text=None):
        self.bar["value"] = percent
        self.value.config(text=text if text else str(percent) + " %")


class MetricsTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent, padding=10)
        self.output = tk.StringVar(value="metrics.csv")
        self.interval = tk.StringVar(value="5")
        self.monitoring = False
        self.stop_event = None
        self.q = queue.Queue()

        if metric_data is None:
            ttk.Label(
                self,
                text="System metrics unavailable: " + str(metrics_err),
            ).pack(anchor="w")
            return

        files = ttk.LabelFrame(self, text="CSV output", padding=8)
        files.pack(fill="x")
        files.columnconfigure(1, weight=1)
        ttk.Label(files, text="Save readings to").grid(row=0, column=0, sticky="w")
        ttk.Entry(files, textvariable=self.output).grid(
            row=0, column=1, sticky="ew", padx=4
        )
        ttk.Button(files, text="Browse", command=self.pick_output).grid(row=0, column=2)
        ttk.Label(files, text="Interval (seconds)").grid(
            row=1, column=0, sticky="w", pady=(6, 0)
        )
        ttk.Entry(files, textvariable=self.interval, width=10).grid(
            row=1, column=1, sticky="w", padx=4, pady=(6, 0)
        )

        readings = ttk.LabelFrame(self, text="Live readings", padding=8)
        readings.pack(fill="x", pady=8)
        self.cpu = Meter(readings, "CPU")
        self.cpu.pack(anchor="w", pady=2)
        self.ram = Meter(readings, "RAM")
        self.ram.pack(anchor="w", pady=2)
        self.disk = Meter(readings, "Disk")
        self.disk.pack(anchor="w", pady=2)
        network = ttk.Frame(readings)
        network.pack(anchor="w", pady=(6, 0))
        self.sent = ttk.Label(network, text="Sent: -")
        self.sent.pack(side="left", padx=(0, 20))
        self.received = ttk.Label(network, text="Received: -")
        self.received.pack(side="left")

        controls = ttk.Frame(self)
        controls.pack(fill="x", pady=(0, 8))
        self.start_btn = ttk.Button(
            controls, text="Start monitoring", command=self.start_monitoring
        )
        self.start_btn.pack(side="left")
        self.stop_btn = ttk.Button(
            controls, text="Stop", command=self.stop_monitoring, state="disabled"
        )
        self.stop_btn.pack(side="left", padx=6)
        self.status = ttk.Label(controls, text="Not monitoring")
        self.status.pack(side="left", padx=8)

        self.log = LogBox(self, height=8)
        self.log.pack(fill="both", expand=True)
        self.after(200, self.drain)

    def pick_output(self):
        path = filedialog.asksaveasfilename(
            title="Save readings as",
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv")],
        )
        if path:
            self.output.set(path)

    def start_monitoring(self):
        if self.monitoring:
            return
        path = self.output.get().strip()
        if not path:
            messagebox.showerror("Output missing", "Choose a CSV output file.")
            return
        try:
            interval = float(self.interval.get())
            if interval < 1:
                raise ValueError
        except ValueError:
            messagebox.showerror(
                "Invalid interval",
                "The interval must be a number of at least 1 second.",
            )
            return

        try:
            needs_header = not os.path.exists(path) or os.path.getsize(path) == 0
            with open(path, "a", newline="", encoding="utf-8") as output:
                if needs_header:
                    csv.writer(output).writerow(metric_data.CSV_COLUMNS)
        except OSError as e:
            messagebox.showerror("Cannot write CSV", str(e))
            return

        self.monitoring = True
        self.stop_event = threading.Event()
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.status.config(text="Monitoring")
        threading.Thread(
            target=self._monitor, args=(path, interval, self.stop_event), daemon=True
        ).start()

    def _monitor(self, path, interval, stop_event):
        try:
            while not stop_event.is_set():
                reading = metric_data.sample_metrics()
                with open(path, "a", newline="", encoding="utf-8") as output:
                    csv.writer(output).writerow(
                        reading[column] for column in metric_data.CSV_COLUMNS
                    )
                self.q.put(("reading", reading))
                stop_event.wait(max(0, interval - 1))
        except Exception as e:
            self.q.put(("error", str(e)))
        finally:
            self.q.put(("stopped", None))

    def stop_monitoring(self):
        if self.stop_event and not self.stop_event.is_set():
            self.stop_event.set()
            self.stop_btn.config(state="disabled")
            if self.status.cget("text") != "Monitoring failed":
                self.status.config(text="Stopping")

    def drain(self):
        while True:
            try:
                event, payload = self.q.get_nowait()
            except queue.Empty:
                break
            if event == "reading":
                self.cpu.set(payload["cpu_%"])
                self.ram.set(payload["ram_%"])
                self.disk.set(payload["disk_%"])
                self.sent.config(text="Sent: " + str(payload["sent_MB"]) + " MB")
                self.received.config(
                    text="Received: " + str(payload["recv_MB"]) + " MB"
                )
                self.log.add(
                    "{}  CPU {}%  RAM {}%  Disk {}%".format(
                        payload["time"],
                        payload["cpu_%"],
                        payload["ram_%"],
                        payload["disk_%"],
                    )
                )
            elif event == "error":
                self.log.add("Monitoring failed: " + payload)
                self.status.config(text="Monitoring failed")
                self.stop_monitoring()
            elif event == "stopped":
                self.monitoring = False
                self.stop_event = None
                self.start_btn.config(state="normal")
                self.stop_btn.config(state="disabled")
                if self.status.cget("text") != "Monitoring failed":
                    self.status.config(text="Not monitoring")
        self.after(200, self.drain)


NONE = "(manual only)"


class SlideDialog(tk.Toplevel):
    def __init__(self, parent, layouts):
        super().__init__(parent)
        self.title("Add slide")
        self.transient(parent.winfo_toplevel())
        self.result = None

        frame = ttk.Frame(self, padding=10)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text="Layout").grid(row=0, column=0, sticky="w", pady=3)
        self.layout = ttk.Combobox(frame, values=layouts, state="readonly", width=34)
        self.layout.current(1 if len(layouts) > 1 else 0)
        self.layout.grid(row=0, column=1, sticky="w")

        ttk.Label(frame, text="Title").grid(row=1, column=0, sticky="w", pady=3)
        self.title_entry = ttk.Entry(frame, width=40)
        self.title_entry.grid(row=1, column=1, sticky="w")

        ttk.Label(frame, text="Bullets\n(one per line)").grid(
            row=2, column=0, sticky="nw", pady=3
        )
        self.bullets = tk.Text(frame, width=40, height=6)
        self.bullets.grid(row=2, column=1, sticky="w")

        ttk.Label(frame, text="Speaker notes").grid(
            row=3, column=0, sticky="nw", pady=3
        )
        self.notes = tk.Text(frame, width=40, height=3)
        self.notes.grid(row=3, column=1, sticky="w")

        buttons = ttk.Frame(frame)
        buttons.grid(row=4, column=0, columnspan=2, sticky="e", pady=(8, 0))
        ttk.Button(buttons, text="Add", command=self.ok).pack(side="left", padx=4)
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="left")

        self.grab_set()
        self.title_entry.focus()
        self.wait_window()

    def ok(self):
        slide = {"layout": int(self.layout.get().split(":")[0])}
        title = self.title_entry.get().strip()
        bullets = [
            l.strip() for l in self.bullets.get("1.0", "end").splitlines() if l.strip()
        ]
        notes = self.notes.get("1.0", "end").strip()
        if title:
            slide["title"] = title
        if bullets:
            slide["bullets"] = bullets
        if notes:
            slide["notes"] = notes
        self.result = slide
        self.destroy()


class PptTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent, padding=10)
        self.template = tk.StringVar()
        self.output = tk.StringVar(value="result.pptx")
        self.preset = tk.StringVar(value="corporate_blue")
        self.heading = tk.StringVar()
        self.body = tk.StringVar()
        self.title_color = tk.StringVar()
        self.body_color = tk.StringVar()
        self.bg = tk.StringVar()
        self.title_size = tk.StringVar()
        self.body_size = tk.StringVar()
        self.slides = []

        self.build()
        for v in (
            self.preset,
            self.heading,
            self.body,
            self.title_color,
            self.body_color,
            self.bg,
        ):
            v.trace_add("write", self.update_preview)
        self.update_preview()

    def build(self):
        self.columnconfigure(0, weight=1)
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, weight=1)
        left = ttk.Frame(self)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        right = ttk.Frame(self)
        right.grid(row=0, column=1, sticky="nsew")

        # 1. files
        f = ttk.LabelFrame(left, text="1. Files", padding=8)
        f.pack(fill="x")
        f.columnconfigure(1, weight=1)
        ttk.Label(f, text="Template").grid(row=0, column=0, sticky="w", pady=2)
        ttk.Entry(f, textvariable=self.template).grid(
            row=0, column=1, sticky="ew", padx=4
        )
        ttk.Button(f, text="Browse", command=self.pick_template).grid(row=0, column=2)
        ttk.Label(f, text="Save as").grid(row=1, column=0, sticky="w", pady=2)
        ttk.Entry(f, textvariable=self.output).grid(
            row=1, column=1, sticky="ew", padx=4
        )
        ttk.Button(f, text="Browse", command=self.pick_output).grid(row=1, column=2)

        # 2. look
        lk = ttk.LabelFrame(left, text="2. Look", padding=8)
        lk.pack(fill="x", pady=8)
        ttk.Label(lk, text="Ready-made theme").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            lk,
            textvariable=self.preset,
            state="readonly",
            width=20,
            values=[NONE] + list(ppt_tool.THEMES),
        ).grid(row=0, column=1, columnspan=3, sticky="w")

        self.preview = tk.Frame(lk, bd=1, relief="solid")
        self.preview.grid(row=1, column=0, columnspan=4, sticky="ew", pady=6)
        self.pv_title = tk.Label(self.preview, text="Sample title")
        self.pv_title.pack(anchor="w", padx=10, pady=(8, 0))
        self.pv_body = tk.Label(self.preview, text="Body text looks like this")
        self.pv_body.pack(anchor="w", padx=10, pady=(0, 8))

        ttk.Label(
            lk,
            text="Manual changes (leave blank to keep the theme's value)",
            foreground="#555555",
        ).grid(row=2, column=0, columnspan=4, sticky="w", pady=(4, 2))

        ttk.Label(lk, text="Heading font").grid(row=3, column=0, sticky="w", pady=2)
        ttk.Combobox(lk, textvariable=self.heading, values=FONTS, width=18).grid(
            row=3, column=1, columnspan=3, sticky="w"
        )
        ttk.Label(lk, text="Body font").grid(row=4, column=0, sticky="w", pady=2)
        ttk.Combobox(lk, textvariable=self.body, values=FONTS, width=18).grid(
            row=4, column=1, columnspan=3, sticky="w"
        )
        self.color_row(lk, 5, "Title color", self.title_color)
        self.color_row(lk, 6, "Body color", self.body_color)
        self.color_row(lk, 7, "Background", self.bg)

        ttk.Label(lk, text="Title size").grid(row=8, column=0, sticky="w", pady=2)
        ttk.Entry(lk, textvariable=self.title_size, width=6).grid(
            row=8, column=1, sticky="w"
        )
        ttk.Label(lk, text="Body size").grid(row=8, column=2, sticky="e")
        ttk.Entry(lk, textvariable=self.body_size, width=6).grid(
            row=8, column=3, sticky="w", padx=4
        )

        # 3. text replacement
        r = ttk.LabelFrame(right, text="3. Text to fill in", padding=8)
        r.pack(fill="both", expand=True)
        ttk.Label(
            r,
            text="One per line, like:  company = ABC\n"
            "(fills {{company}} in the template)",
            foreground="#555555",
        ).pack(anchor="w")
        self.replace_text = tk.Text(r, height=7, width=40)
        self.replace_text.pack(fill="both", expand=True, pady=4)

        # 4. new slides
        s = ttk.LabelFrame(right, text="4. New slides (optional)", padding=8)
        s.pack(fill="both", expand=True, pady=(8, 0))
        self.slide_list = tk.Listbox(s, height=6)
        self.slide_list.pack(fill="both", expand=True)
        row = ttk.Frame(s)
        row.pack(fill="x", pady=(4, 0))
        ttk.Button(row, text="Add slide...", command=self.add_slide).pack(side="left")
        ttk.Button(row, text="Remove selected", command=self.remove_slide).pack(
            side="left", padx=6
        )

        # bottom
        self.status = ttk.Label(self, text="")
        self.status.grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Button(self, text="Create presentation", command=self.make_ppt).grid(
            row=1, column=1, sticky="e", pady=(8, 0)
        )

    def color_row(self, parent, row, label, var):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=2)
        ttk.Entry(parent, textvariable=var, width=10).grid(
            row=row, column=1, sticky="w"
        )
        swatch = tk.Label(parent, width=3, relief="solid", bd=1, bg="#eeeeee")
        swatch.grid(row=row, column=2, padx=4)

        def pick():
            start = norm_hex(var.get()) if var.get().strip() else None
            chosen = colorchooser.askcolor(color=start, title=label)[1]
            if chosen:
                var.set(chosen.upper())

        def refresh(*args):
            value = norm_hex(var.get()) if var.get().strip() else None
            swatch.config(bg=value if value else "#eeeeee")

        var.trace_add("write", refresh)
        ttk.Button(parent, text="Pick", width=5, command=pick).grid(row=row, column=3)

    # ---- file pickers
    def pick_template(self):
        path = filedialog.askopenfilename(
            title="Choose template", filetypes=[("PowerPoint", "*.pptx *.potx")]
        )
        if path:
            self.template.set(path)

    def pick_output(self):
        path = filedialog.asksaveasfilename(
            title="Save as",
            defaultextension=".pptx",
            filetypes=[("PowerPoint", "*.pptx")],
        )
        if path:
            self.output.set(path)

    # ---- slides list
    def get_layouts(self):
        # read the real layout names from the template so nobody has to guess numbers
        try:
            prs = ppt_tool.Presentation(self.template.get().strip())
            return [str(i) + ": " + l.name for i, l in enumerate(prs.slide_layouts)]
        except Exception:
            return [str(i) for i in range(8)]

    def add_slide(self):
        dialog = SlideDialog(self, self.get_layouts())
        if dialog.result:
            self.slides.append(dialog.result)
            title = dialog.result.get("title", "(no title)")
            self.slide_list.insert(
                "end", "Layout " + str(dialog.result["layout"]) + " - " + title
            )

    def remove_slide(self):
        picked = self.slide_list.curselection()
        if picked:
            self.slides.pop(picked[0])
            self.slide_list.delete(picked[0])

    # ---- build the config dictionary from the form
    def collect(self, strict):
        config = {}
        if self.preset.get() != NONE:
            config["preset"] = self.preset.get()

        theme = {}
        if self.heading.get().strip():
            theme["heading_font"] = self.heading.get().strip()
        if self.body.get().strip():
            theme["body_font"] = self.body.get().strip()
        if theme:
            config["theme"] = theme

        style = {}
        for key, var, label in (
            ("title_color", self.title_color, "Title color"),
            ("body_color", self.body_color, "Body color"),
        ):
            if not var.get().strip():
                continue
            value = norm_hex(var.get())
            if value is None:
                if strict:
                    raise ValueError(label + " must look like #1F4E79")
                continue
            style[key] = value
        for key, var, label in (
            ("title_size", self.title_size, "Title size"),
            ("body_size", self.body_size, "Body size"),
        ):
            if not var.get().strip():
                continue
            try:
                style[key] = int(var.get())
            except ValueError:
                if strict:
                    raise ValueError(label + " must be a number")
        if style:
            config["text_style"] = style

        if self.bg.get().strip():
            value = norm_hex(self.bg.get())
            if value is None:
                if strict:
                    raise ValueError("Background must look like #F7F7F2")
            else:
                config["background_color"] = value

        if strict:
            replacements = {}
            for line in self.replace_text.get("1.0", "end").splitlines():
                if "=" not in line:
                    continue
                key, value = line.split("=", 1)
                key = key.strip()
                if not key:
                    continue
                if not key.startswith("{{"):
                    key = "{{" + key + "}}"
                replacements[key] = value.strip()
            if replacements:
                config["replacements"] = replacements
            if self.slides:
                config["slides"] = self.slides
        return config

    def update_preview(self, *args):
        try:
            theme, style, bg = ppt_tool.build_settings(self.collect(strict=False))
        except Exception:
            return
        bg = bg or "#FFFFFF"
        self.preview.config(bg=bg)
        self.pv_title.config(
            bg=bg,
            fg=style.get("title_color", "#000000"),
            font=(theme.get("heading_font", "Calibri"), 20, "bold"),
        )
        self.pv_body.config(
            bg=bg,
            fg=style.get("body_color", "#000000"),
            font=(theme.get("body_font", "Calibri"), 12),
        )

    def make_ppt(self):
        template = self.template.get().strip()
        if not template or not os.path.isfile(template):
            messagebox.showerror("Template missing", "Choose a template file first.")
            return
        try:
            config = self.collect(strict=True)
            config["template"] = template
            config["output"] = self.output.get().strip() or "result.pptx"
            path = ppt_tool.run(config)
        except PermissionError:
            messagebox.showerror(
                "Cannot save",
                "The output file is open in PowerPoint. Close it and try again.",
            )
            return
        except Exception as e:
            messagebox.showerror("Could not create the presentation", str(e))
            return

        self.status.config(text="Saved: " + path)
        if messagebox.askyesno("Done", "Saved " + path + "\n\nOpen it now?"):
            open_file(path)


# ---------------------------------------------------------------- tab 2: excel to xml
class ExcelTab(ttk.Frame):
    CHECK_EVERY = 3000  # ms

    def __init__(self, parent):
        super().__init__(parent, padding=10)
        self.file = tk.StringVar()
        self.sheet = tk.StringVar()
        self.xml = tk.StringVar()
        self.folder = tk.StringVar()
        self.out_folder = tk.StringVar()
        self.watching = False
        self.busy = False
        self.q = queue.Queue()

        # convert one file
        one = ttk.LabelFrame(self, text="Convert one file", padding=8)
        one.pack(fill="x")
        one.columnconfigure(1, weight=1)
        ttk.Label(one, text="Excel file").grid(row=0, column=0, sticky="w", pady=2)
        ttk.Entry(one, textvariable=self.file).grid(
            row=0, column=1, sticky="ew", padx=4
        )
        ttk.Button(one, text="Browse", command=self.pick_excel).grid(row=0, column=2)
        ttk.Label(one, text="Sheet").grid(row=1, column=0, sticky="w", pady=2)
        self.sheet_box = ttk.Combobox(one, textvariable=self.sheet, width=24)
        self.sheet_box.grid(row=1, column=1, sticky="w", padx=4)
        ttk.Label(one, text="(blank = first/active sheet)", foreground="#555555").grid(
            row=1, column=2, sticky="w"
        )
        ttk.Label(one, text="Save XML as").grid(row=2, column=0, sticky="w", pady=2)
        ttk.Entry(one, textvariable=self.xml).grid(row=2, column=1, sticky="ew", padx=4)
        ttk.Button(one, text="Browse", command=self.pick_xml).grid(row=2, column=2)
        ttk.Button(one, text="Convert", command=self.convert_one).grid(
            row=3, column=2, sticky="e", pady=(6, 0)
        )

        # watch a folder
        w = ttk.LabelFrame(
            self,
            text="Watch a folder (converts every Excel file dropped into it)",
            padding=8,
        )
        w.pack(fill="x", pady=10)
        w.columnconfigure(1, weight=1)
        ttk.Label(w, text="Folder to watch").grid(row=0, column=0, sticky="w", pady=2)
        ttk.Entry(w, textvariable=self.folder).grid(
            row=0, column=1, sticky="ew", padx=4
        )
        ttk.Button(
            w, text="Browse", command=lambda: self.pick_folder(self.folder)
        ).grid(row=0, column=2)
        ttk.Label(w, text="XML output folder").grid(row=1, column=0, sticky="w", pady=2)
        ttk.Entry(w, textvariable=self.out_folder).grid(
            row=1, column=1, sticky="ew", padx=4
        )
        ttk.Button(
            w, text="Browse", command=lambda: self.pick_folder(self.out_folder)
        ).grid(row=1, column=2)
        ttk.Label(
            w,
            text="(blank = an 'xml_output' folder inside the watched folder)",
            foreground="#555555",
        ).grid(row=2, column=1, sticky="w")
        row = ttk.Frame(w)
        row.grid(row=3, column=0, columnspan=3, sticky="w", pady=(6, 0))
        self.start_btn = ttk.Button(
            row, text="Start watching", command=self.start_watch
        )
        self.start_btn.pack(side="left")
        self.stop_btn = ttk.Button(
            row, text="Stop", command=self.stop_watch, state="disabled"
        )
        self.stop_btn.pack(side="left", padx=6)
        self.state_label = ttk.Label(row, text="not watching")
        self.state_label.pack(side="left", padx=10)

        ttk.Label(self, text="Log").pack(anchor="w")
        self.log = LogBox(self, height=10)
        self.log.pack(fill="both", expand=True)

        self.after(200, self.drain)

    def pick_excel(self):
        path = filedialog.askopenfilename(
            title="Choose Excel file", filetypes=[("Excel", "*.xlsx *.xlsm")]
        )
        if not path:
            return
        self.file.set(path)
        if not self.xml.get():
            self.xml.set(os.path.splitext(path)[0] + ".xml")
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as e:
            self.sheet_box["values"] = []
            self.log.add("Could not load workbook support: " + str(e))
            return
        try:
            wb = load_workbook(path, read_only=True)
            try:
                self.sheet_box["values"] = wb.sheetnames
            finally:
                wb.close()
        except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as e:
            self.sheet_box["values"] = []
            self.log.add("Could not read workbook sheets: " + str(e))

    def pick_xml(self):
        path = filedialog.asksaveasfilename(
            title="Save XML as", defaultextension=".xml", filetypes=[("XML", "*.xml")]
        )
        if path:
            self.xml.set(path)

    def pick_folder(self, variable):
        path = filedialog.askdirectory(title="Choose folder")
        if path:
            variable.set(path)

    @classmethod
    def _convert_file(cls, source, destination, sheet=None):
        if excel_to_xml is None:
            raise RuntimeError("Excel-to-XML converter unavailable: " + str(excel_err))
        converted = excel_to_xml.convert(source, destination, sheet)
        if not converted:
            raise RuntimeError("The selected worksheet is empty; no XML was created.")
        return "Converted " + os.path.basename(source) + " to " + destination

    def convert_one(self):
        if self.busy:
            return
        if excel_to_xml is None:
            messagebox.showerror(
                "Converter unavailable",
                "Excel-to-XML converter unavailable: " + str(excel_err),
            )
            return
        source = self.file.get().strip()
        destination = self.xml.get().strip()
        if not source or not os.path.isfile(source):
            messagebox.showerror(
                "Excel file missing", "Choose an existing Excel file first."
            )
            return
        if not destination:
            destination = os.path.splitext(source)[0] + ".xml"
            self.xml.set(destination)
        self.busy = True
        sheet = self.sheet.get().strip() or None

        def convert():
            try:
                output = self._convert_file(source, destination, sheet)
                self.q.put(("log", output))
            except Exception as e:
                self.q.put(("log", "Conversion failed: " + str(e)))
            finally:
                self.q.put(("convert_done", None))

        threading.Thread(target=convert, daemon=True).start()

    @staticmethod
    def _move_processed(source, folder):
        destination = os.path.join(folder, os.path.basename(source))
        if os.path.exists(destination):
            base, extension = os.path.splitext(destination)
            stamp = datetime.now().strftime("%H%M%S%f")
            destination = base + "_" + stamp + extension
        shutil.move(source, destination)
        return destination

    def start_watch(self):
        if self.watching:
            return
        if excel_to_xml is None:
            messagebox.showerror(
                "Converter unavailable",
                "Excel-to-XML converter unavailable: " + str(excel_err),
            )
            return
        folder = self.folder.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror(
                "Folder missing", "Choose an existing folder to watch."
            )
            return

        out_folder = self.out_folder.get().strip() or os.path.join(folder, "xml_output")
        done_folder = os.path.join(folder, "processed")
        failed_folder = os.path.join(folder, "failed")
        try:
            for path in (out_folder, done_folder, failed_folder):
                os.makedirs(path, exist_ok=True)
        except OSError as e:
            messagebox.showerror("Could not start watching", str(e))
            return

        self.watching = True
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.state_label.config(text="watching")
        threading.Thread(
            target=self._watch_folder,
            args=(folder, out_folder, done_folder, failed_folder),
            daemon=True,
        ).start()

    def _watch_folder(self, folder, out_folder, done_folder, failed_folder):
        try:
            while self.watching:
                for name in os.listdir(folder):
                    if not self.watching:
                        break
                    if name.startswith("~$") or not name.lower().endswith(
                        (".xlsx", ".xlsm")
                    ):
                        continue
                    source = os.path.join(folder, name)
                    if not os.path.isfile(source):
                        continue
                    try:
                        initial_size = os.path.getsize(source)
                        time.sleep(1)
                        if not self.watching or os.path.getsize(source) != initial_size:
                            continue
                        destination = os.path.join(
                            out_folder, os.path.splitext(name)[0] + ".xml"
                        )
                        output = self._convert_file(source, destination)
                        moved_to = self._move_processed(source, done_folder)
                        self.q.put(("log", output + " - moved to " + moved_to))
                    except Exception as e:
                        self.q.put(("log", "Could not convert " + name + ": " + str(e)))
                        try:
                            self._move_processed(source, failed_folder)
                        except OSError as move_error:
                            self.q.put(
                                (
                                    "log",
                                    "Could not move "
                                    + name
                                    + " to failed: "
                                    + str(move_error),
                                )
                            )
                for _ in range(self.CHECK_EVERY // 100):
                    if not self.watching:
                        break
                    time.sleep(0.1)
        except OSError as e:
            self.q.put(("log", "Folder watch failed: " + str(e)))
        finally:
            self.q.put(("watch_stopped", None))

    def stop_watch(self):
        if self.watching:
            self.watching = False
            self.stop_btn.config(state="disabled")
            self.state_label.config(text="stopping")

    def drain(self):
        while True:
            try:
                event, payload = self.q.get_nowait()
            except queue.Empty:
                break
            if event == "log":
                self.log.add(payload)
            elif event == "convert_done":
                self.busy = False
            elif event == "watch_stopped":
                self.watching = False
                self.start_btn.config(state="normal")
                self.stop_btn.config(state="disabled")
                self.state_label.config(text="not watching")
        self.after(200, self.drain)


class MailTab(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent, padding=10)
        self.recipients = tk.StringVar()
        self.subject = tk.StringVar()
        self.files = []

        self.columnconfigure(1, weight=1)
        ttk.Label(self, text="To (separate email addresses with commas)").grid(
            row=0, column=0, sticky="w", pady=4
        )
        recipient_entry = ttk.Entry(self, textvariable=self.recipients)
        recipient_entry.grid(row=0, column=1, columnspan=2, sticky="ew", padx=6)
        ttk.Label(self, text="Subject").grid(row=1, column=0, sticky="w", pady=4)
        subject_entry = ttk.Entry(self, textvariable=self.subject)
        subject_entry.grid(row=1, column=1, columnspan=2, sticky="ew", padx=6)

        files_frame = ttk.LabelFrame(self, text="Attachments", padding=8)
        files_frame.grid(row=2, column=0, columnspan=3, sticky="nsew", pady=8)
        files_frame.columnconfigure(0, weight=1)
        files_frame.rowconfigure(0, weight=1)
        self.rowconfigure(2, weight=1)
        self.file_list = tk.Listbox(files_frame, height=8, selectmode="extended")
        self.file_list.grid(row=0, column=0, rowspan=2, sticky="nsew")
        ttk.Button(files_frame, text="Add files...", command=self.add_files).grid(
            row=0, column=1, sticky="ew", padx=(8, 0), pady=(0, 4)
        )
        ttk.Button(
            files_frame, text="Remove selected", command=self.remove_files
        ).grid(row=1, column=1, sticky="ew", padx=(8, 0), pady=(4, 0))

        ttk.Label(self, text="Generated command").grid(
            row=3, column=0, columnspan=3, sticky="w", pady=(4, 2)
        )
        self.command = tk.Text(self, height=4, wrap="word", state="disabled")
        self.command.grid(row=4, column=0, columnspan=3, sticky="ew")
        self.status = ttk.Label(
            self,
            text="Opening a new console runs this command; cmd_mail.py will prompt for your email and app password.",
            wraplength=760,
        )
        self.status.grid(row=5, column=0, columnspan=3, sticky="w", pady=6)
        self.open_button = ttk.Button(
            self, text="Open console and run", command=self.open_command
        )
        self.open_button.grid(row=6, column=2, sticky="e", pady=(4, 0))

        self.recipients.trace_add("write", self.update_command)
        self.subject.trace_add("write", self.update_command)
        self.update_command()

    def add_files(self):
        paths = filedialog.askopenfilenames(title="Choose files to attach")
        for path in paths:
            if path not in self.files:
                self.files.append(path)
                self.file_list.insert("end", path)
        self.update_command()

    def remove_files(self):
        for index in reversed(self.file_list.curselection()):
            self.file_list.delete(index)
            del self.files[index]
        self.update_command()

    def build_command(self):
        script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cmd_mail.py")
        if not os.path.isfile(script):
            raise FileNotFoundError("Could not find cmd_mail.py beside app.py.")
        python_executable = sys.executable
        if os.path.basename(python_executable).lower() == "pythonw.exe":
            console_python = os.path.join(os.path.dirname(python_executable), "python.exe")
            if not os.path.isfile(console_python):
                raise FileNotFoundError(
                    "Could not find python.exe beside pythonw.exe; a console-enabled Python is required."
                )
            python_executable = console_python

        recipients = [
            address.strip()
            for address in re.split(r"[,;]", self.recipients.get())
            if address.strip()
        ]
        if not recipients or not self.files:
            return ""

        arguments = [python_executable, script, "--to", *recipients]
        subject = self.subject.get().strip()
        if subject:
            arguments.extend(["--subject", subject])
        arguments.extend(["--files", *self.files])
        arguments.append("--pause")
        return arguments

    def update_command(self, *_args):
        try:
            arguments = self.build_command()
            command = subprocess.list2cmdline(arguments) if arguments else ""
            error = None
        except (OSError, ValueError) as e:
            arguments = []
            command = ""
            error = str(e)

        self.command.config(state="normal")
        self.command.delete("1.0", "end")
        self.command.insert("1.0", command)
        self.command.config(state="disabled")
        self.open_button.config(state="normal" if command else "disabled")
        if error:
            self.status.config(text=error)
        elif not self.files:
            self.status.config(text="Add at least one attachment to generate the command.")
        elif not self.recipients.get().strip():
            self.status.config(text="Enter at least one recipient to generate the command.")
        else:
            self.status.config(
                text="Opening a new console runs this command; cmd_mail.py will prompt for your email and app password."
            )

    def open_command(self):
        try:
            arguments = self.build_command()
            if not arguments:
                self.update_command()
                return
            subprocess.Popen(
                arguments,
                cwd=os.path.dirname(os.path.abspath(__file__)),
                creationflags=subprocess.CREATE_NEW_CONSOLE,
            )
        except (FileNotFoundError, OSError, ValueError) as e:
            messagebox.showerror("Could not open console", str(e))


def main():
    root = tk.Tk()
    root.title("Desktop Tools")
    root.geometry("1000x700")

    tabs = ttk.Notebook(root)
    tabs.pack(fill="both", expand=True)
    if ppt_tool is None:
        ppt_frame = ttk.Frame(tabs, padding=12)
        ttk.Label(ppt_frame, text="PowerPoint tools unavailable: " + str(ppt_err)).pack(
            anchor="w"
        )
        tabs.add(ppt_frame, text="PowerPoint")
    else:
        tabs.add(PptTab(tabs), text="PowerPoint")
    tabs.add(ExcelTab(tabs), text="Excel to XML")
    tabs.add(MetricsTab(tabs), text="System Metrics")
    tabs.add(MailTab(tabs), text="Send Mail")

    root.mainloop()


if __name__ == "__main__":
    main()
